import openpyxl
from docx import Document
from PyPDF2 import PdfReader
from flask import Flask, make_response
from flask import request, jsonify
from flask_cors import CORS
from werkzeug.utils import secure_filename
import os
from functools import reduce
app = Flask(__name__)

cors = CORS(app, resources={r'*': {'origins': '*'}})

# handeling PDF


def pdfProcessing(path):
    reader = PdfReader(path)
    number_of_pages = len(reader.pages)
    text = ""
    for i in range(number_of_pages):
        page = reader.pages[i]
        text += page.extract_text()
    return text


# handleling Doc


def docProcessing(path):
    doc_object = open(path, "rb")
    # creating word reader object
    doc_reader = Document(doc_object)
    data = ""
    for p in doc_reader.paragraphs:
        data += p.text+"\n"
    return data


# handleling xle
def xlsProcessing(path):
    workbook = openpyxl.load_workbook(path)
    worksheet = workbook['Sheet1']
    extracted_text = []
    for row in worksheet.iter_rows(values_only=True):
        cellText = ""
        for cell_value in row:
            print(str(cell_value))
            cellText = cellText+" "+str(cell_value)
        extracted_text.append(cellText)
    extracted_text = '. \n'.join(extracted_text)
    # Close the workbook when done
    workbook.close()
    return extracted_text


@app.route("/api/serverprocess", methods=['POST'])
def serverprocess():
    res_files = request.files
    # saving files to server
    os.makedirs(os.path.join(app.instance_path, 'htmlfi'), exist_ok=True)
    for k, v in res_files.items():
        v.save(os.path.join(app.instance_path,
               'htmlfi', secure_filename(v.filename)))

    # extracting content of uploaded files
    response = []
    for k, v in res_files.items():
        extention = k.split('.')
        extention = extention[len(extention)-1]
        if extention == 'pdf':
            text = pdfProcessing(os.path.join(
                app.instance_path, 'htmlfi', secure_filename(k)))
            obj = {"filename": k, "content": text}
            response.append(obj)
        elif extention == 'docx':
            text = docProcessing(os.path.join(
                app.instance_path, 'htmlfi', secure_filename(k)))
            obj = {"filename": k, "content": text}
            response.append(obj)
        elif extention == 'xlsx':
            text = xlsProcessing(os.path.join(
                app.instance_path, 'htmlfi', secure_filename(k)))
            obj = {"filename": k, "content": text}
            response.append(obj)

    return jsonify(response)


if __name__ == '__main__':
    app.run(debug=True, port=8080)
